from playwright.sync_api import sync_playwright
import time
import os
import argparse
import re
from pathlib import Path
import sys
import openpyxl
from openpyxl.cell.cell import MergedCell

# Configuration
ROOT_DIR = Path(__file__).resolve().parent
DEFAULT_FRONTEND_URL = os.getenv("FRONTEND_URL", "https://tmrtools.com/")

DEFAULT_INPUT_COLUMN_CANDIDATES = [
    "Singlish",
    "Input",
    "Singlish Input",
    "Test Input",
    "Source",
    "Sentence",
    "Text",
]

DEFAULT_EXPECTED_COLUMN_CANDIDATES = [
    "Sinhala",
    "Expected_Output",
    "Expected Output",
    "Expected output",
    "Expected",
    "Expected Sinhala",
]

DEFAULT_ACTUAL_COLUMN_CANDIDATES = [
    "Actual_Output",
    "Actual Output",
    "Actual output",
    "Actual",
]

DEFAULT_STATUS_COLUMN_CANDIDATES = [
    "Status",
    "Result",
    "Pass/Fail",
    "Pass Fail",
]

DEFAULT_WAIT_MS = 10000
DEFAULT_RETRIES = 30
DEFAULT_RETRY_WAIT_MS = 2000
DEFAULT_TYPE_DELAY_MS = 30
DEFAULT_TIMEOUT_MS = 60000
DEFAULT_SLOW_MO_MS = 0


def _configure_stdout():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
    except Exception:
        pass


def _resolve_path(p: str | None) -> str | None:
    if not p:
        return None
    path = Path(p)
    if path.is_absolute():
        return str(path)
    return str((ROOT_DIR / path).resolve())


def _normalize_header(value) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def _header_values(ws, row_index: int) -> list:
    max_col = max(1, int(ws.max_column or 1))
    return [ws.cell(row=row_index, column=c).value for c in range(1, max_col + 1)]


def _find_header_row(ws, max_scan_rows: int) -> int:
    expected_tokens = {_normalize_header(v) for v in DEFAULT_EXPECTED_COLUMN_CANDIDATES}

    scan_limit = max(1, min(int(max_scan_rows), int(ws.max_row or 1)))
    for r in range(1, scan_limit + 1):
        values = _header_values(ws, r)
        texts = [v for v in values if isinstance(v, str) and v.strip()]
        norms = {_normalize_header(v) for v in texts}

        if "input" in norms and "expectedoutput" in norms:
            return r

        if "input" in norms and (norms & expected_tokens):
            return r

    return 1


def _merged_top_left_cell(ws, row: int, col: int):
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return cell

    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col)

    return ws.cell(row=row, column=col)


def _is_top_left_of_merged_cell(ws, row: int, col: int) -> bool:
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return True

    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return rng.min_row == row and rng.min_col == col

    return True


def _set_cell_value(ws, row: int, col: int, value):
    cell = _merged_top_left_cell(ws, row, col)
    cell.value = value


def _find_column_index(header_values: list, requested_name: str | None, candidates: list[str]) -> int | None:
    indexed = []
    for i, v in enumerate(header_values, start=1):
        if v is None:
            continue
        indexed.append((i, str(v)))

    norm_to_index: dict[str, int] = {}
    for i, v in indexed:
        n = _normalize_header(v)
        if n and n not in norm_to_index:
            norm_to_index[n] = i

    def match(name: str) -> int | None:
        n = _normalize_header(name)
        if not n:
            return None

        if n in norm_to_index:
            return norm_to_index[n]

        for i, v in indexed:
            header_norm = _normalize_header(v)
            if n in header_norm or header_norm in n:
                return i

        return None

    if requested_name:
        found = match(requested_name)
        if found:
            return found

    for c in candidates:
        found = match(c)
        if found:
            return found

    return None


def _last_header_col(header_values: list) -> int:
    last = 0
    for i, v in enumerate(header_values, start=1):
        if v is None:
            continue
        if isinstance(v, str) and not v.strip():
            continue
        last = i
    return last


def _ensure_column(ws, header_row: int, header_values: list, desired_name: str) -> int:
    found = _find_column_index(header_values, desired_name, [])
    if found:
        return found

    col = _last_header_col(header_values) + 1
    ws.cell(row=header_row, column=col).value = desired_name

    if col <= len(header_values):
        header_values[col - 1] = desired_name
    else:
        while len(header_values) < col - 1:
            header_values.append(None)
        header_values.append(desired_name)

    return col


def _dismiss_overlays(page):
    candidates = [
        ("button", re.compile(r"^(Accept|I Agree|Agree|OK|Got it)$", re.IGNORECASE)),
        ("button", re.compile(r"^(Accept all|Accept All)$", re.IGNORECASE)),
    ]

    for role, name in candidates:
        try:
            btn = page.get_by_role(role, name=name).first
            if btn.is_visible():
                btn.click(timeout=2000)
                page.wait_for_timeout(500)
        except Exception:
            pass


def _find_input_locator(page, timeout_ms: int):
    deadline = time.time() + (max(1, timeout_ms) / 1000)

    while time.time() < deadline:
        _dismiss_overlays(page)

        candidates = [
            page.locator('textarea[placeholder*="Singlish"]').first,
            page.locator('textarea[placeholder*="English"]').first,
            page.locator('textarea').first,
            page.locator('[contenteditable="true"]').first,
            page.get_by_role("textbox").first,
        ]

        for loc in candidates:
            try:
                if loc.count() > 0 and loc.is_visible():
                    return loc
            except Exception:
                pass

        page.wait_for_timeout(500)

    raise RuntimeError("Could not find TMRTools input field.")


def _clear_and_type(page, locator, text: str, type_delay_ms: int):
    try:
        locator.click(timeout=3000)
    except Exception:
        pass

    try:
        locator.fill("")
    except Exception:
        try:
            page.keyboard.press("Control+A")
            page.keyboard.press("Backspace")
        except Exception:
            pass

    try:
        if type_delay_ms and int(type_delay_ms) > 0:
            locator.type(text, delay=int(type_delay_ms))
        else:
            locator.fill(text)
    except Exception:
        page.keyboard.insert_text(text)


def _click_translate_if_exists(page):
    button_names = [
        r"^Translate$",
        r"^Transliterate$",
        r"^Convert$",
        r"^Submit$",
        r"^Generate$",
    ]

    for name in button_names:
        try:
            btn = page.get_by_role("button", name=re.compile(name, re.IGNORECASE)).first
            if btn.count() > 0 and btn.is_visible():
                btn.click(timeout=3000)
                return True
        except Exception:
            pass

    return False


def _read_sinhala_output(page, input_text: str) -> str:
    # This tries to find visible Sinhala text on the page, excluding input and fixed UI labels.
    try:
        return page.evaluate(
            """
            (inputText) => {
                const isVisible = (el) => {
                    const style = window.getComputedStyle(el);
                    return style
                        && style.visibility !== 'hidden'
                        && style.display !== 'none'
                        && (el.offsetWidth || el.offsetHeight || el.getClientRects().length);
                };

                const getText = (el) => {
                    if ('value' in el) return (el.value || '').trim();
                    return (el.innerText || el.textContent || '').trim();
                };

                const uiLabels = [
                    'Singlish & English to Sinhala Translator',
                    'Easily convert Singlish or English text into Sinhala script.',
                    'Input Source',
                    'Text',
                    'Upload File',
                    'Type Singlish',
                    'Sinhala Output',
                    'Translated Sinhala text will appear here',
                    'Use toolbar to format',
                    'All rights reserved'
                ];

                const nodes = Array.from(document.querySelectorAll(
                    'textarea, input, [contenteditable="true"], [role="textbox"], div, p, span'
                ));

                const candidates = nodes
                    .filter(isVisible)
                    .map(getText)
                    .filter(t => t && t.length > 0)
                    .filter(t => t !== inputText.trim())
                    .filter(t => !uiLabels.some(label => t.includes(label)))
                    .filter(t => /[\\u0D80-\\u0DFF]/.test(t));

                if (candidates.length === 0) return '';

                candidates.sort((a, b) => b.length - a.length);
                return candidates[0].trim();
            }
            """,
            input_text,
        ).strip()
    except Exception:
        return ""


def _parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", required=True)
    parser.add_argument("--sheet", default=None)
    parser.add_argument("--header-row", type=int, default=0)
    parser.add_argument("--max-header-scan-rows", type=int, default=30)
    parser.add_argument("--input-col", default=None)
    parser.add_argument("--expected-col", default=None)
    parser.add_argument("--actual-col", default=None)
    parser.add_argument("--status-col", default=None)
    parser.add_argument("--url", default=DEFAULT_FRONTEND_URL)
    parser.add_argument("--output", default=None)
    parser.add_argument("--save-every", type=int, default=1)
    parser.add_argument("--headless", action="store_true", default=False)
    parser.add_argument("--wait-ms", type=int, default=DEFAULT_WAIT_MS)
    parser.add_argument("--retries", type=int, default=DEFAULT_RETRIES)
    parser.add_argument("--retry-wait-ms", type=int, default=DEFAULT_RETRY_WAIT_MS)
    parser.add_argument("--type-delay-ms", type=int, default=DEFAULT_TYPE_DELAY_MS)
    parser.add_argument("--timeout-ms", type=int, default=DEFAULT_TIMEOUT_MS)
    parser.add_argument("--slow-mo-ms", type=int, default=DEFAULT_SLOW_MO_MS)
    return parser.parse_args()


def run_test():
    _configure_stdout()
    args = _parse_args()
    args.excel = _resolve_path(args.excel)
    args.output = _resolve_path(args.output) if args.output else args.excel

    if not args.excel or not os.path.exists(args.excel):
        print(f"Error: File '{args.excel}' not found.")
        return

    try:
        wb = openpyxl.load_workbook(args.excel)
    except Exception as e:
        print(f"Error reading Excel file: {e}")
        return

    if args.sheet and args.sheet in wb.sheetnames:
        ws = wb[args.sheet]
    else:
        ws = wb.active

    header_row = int(args.header_row or 0)
    if header_row <= 0:
        header_row = _find_header_row(ws, int(args.max_header_scan_rows))

    header_values = _header_values(ws, header_row)

    input_col_idx = _find_column_index(header_values, args.input_col, DEFAULT_INPUT_COLUMN_CANDIDATES)
    expected_col_idx = _find_column_index(header_values, args.expected_col, DEFAULT_EXPECTED_COLUMN_CANDIDATES)

    if not input_col_idx:
        printable = [str(v) if v is not None else "" for v in header_values]
        print("Error: Could not resolve input column.")
        print(f"Header row: {header_row}")
        print(f"Available columns: {printable}")
        return

    actual_col_name = args.actual_col or "Actual output"
    status_col_name = args.status_col or "Status"

    actual_col_idx = _find_column_index(header_values, args.actual_col, DEFAULT_ACTUAL_COLUMN_CANDIDATES)
    status_col_idx = _find_column_index(header_values, args.status_col, DEFAULT_STATUS_COLUMN_CANDIDATES)

    actual_col_idx = actual_col_idx or _ensure_column(ws, header_row, header_values, actual_col_name)
    status_col_idx = status_col_idx or _ensure_column(ws, header_row, header_values, status_col_name)

    rows_total = max(0, int(ws.max_row or 0) - header_row)
    print(f"Starting TMRTools test with {rows_total} rows...")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=args.headless, slow_mo=max(0, int(args.slow_mo_ms)))
        page = browser.new_page()
        page.set_default_timeout(max(1000, int(args.timeout_ms)))

        processed = 0

        for row_index in range(header_row + 1, int(ws.max_row or 0) + 1):
            if not _is_top_left_of_merged_cell(ws, row_index, input_col_idx):
                continue

            input_cell = _merged_top_left_cell(ws, row_index, input_col_idx)
            input_value = input_cell.value
            singlish_input = str(input_value).strip() if input_value is not None else ""

            if not singlish_input:
                continue

            expected_value = (
                _merged_top_left_cell(ws, row_index, expected_col_idx).value if expected_col_idx else None
            )
            expected_sinhala = str(expected_value).strip() if expected_value is not None else ""

            print(f"Testing [Row {row_index}]: {singlish_input}")

            try:
                page.goto(args.url, wait_until="domcontentloaded")
                try:
                    page.wait_for_load_state("networkidle", timeout=max(1000, int(args.timeout_ms)))
                except Exception:
                    pass

                page.wait_for_timeout(1500)

                input_locator = _find_input_locator(page, int(args.timeout_ms))
                _clear_and_type(page, input_locator, singlish_input, int(args.type_delay_ms))

                _click_translate_if_exists(page)

                page.wait_for_timeout(max(0, int(args.wait_ms)))

                actual_output = ""

                for _ in range(max(1, int(args.retries))):
                    current = _read_sinhala_output(page, singlish_input)
                    if current:
                        actual_output = current
                        break
                    page.wait_for_timeout(max(0, int(args.retry_wait_ms)))

                if not actual_output:
                    actual_output = "No output generated within waiting time"

                _set_cell_value(ws, row_index, actual_col_idx, actual_output)

                if expected_sinhala:
                    status = "PASS" if actual_output == expected_sinhala else "FAIL"
                else:
                    status = "COLLECTED"

                _set_cell_value(ws, row_index, status_col_idx, status)

                print(f"  -> {status}")
                processed += 1

                if args.save_every and int(args.save_every) > 0 and processed % int(args.save_every) == 0:
                    wb.save(args.output)

            except Exception as e:
                print(f"Error in UI interaction: {e}")
                try:
                    _set_cell_value(ws, row_index, actual_col_idx, f"UI Error: {e}")
                    _set_cell_value(ws, row_index, status_col_idx, "UI Error")
                except Exception:
                    pass

                if args.save_every and int(args.save_every) > 0:
                    try:
                        wb.save(args.output)
                    except Exception:
                        pass

        browser.close()

    try:
        wb.save(args.output)
    except Exception as e:
        print(f"Error saving output file '{args.output}': {e}")
        return

    print(f"Test completed. Results saved to {args.output}")


if __name__ == "__main__":
    run_test()